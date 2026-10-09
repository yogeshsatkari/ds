import re
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

FONT_NAME = "Calibri"
BLACK_COLOR = RGBColor(0, 0, 0)


def apply_cell_padding(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins (padding) for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def apply_inline_formatting(paragraph, text: str, default_size_pt: float = 11):
    """Parses inline **bold** and *italic* formatting."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)
            
        run.font.name = FONT_NAME
        run.font.size = Pt(default_size_pt)
        run.font.color.rgb = BLACK_COLOR


def render_markdown_table(doc: DocumentType, table_lines: list[str]):
    """Parses and renders a GitHub-flavored Markdown table with styling."""
    rows = []
    for line in table_lines:
        clean = line.strip()
        # Skip markdown table divider line (| --- | :---: |)
        if not clean or all(c in "-: |" for c in clean):
            continue
        # Split cells
        cells = [c.strip() for c in clean.strip("|").split("|")]
        rows.append(cells)

    if not rows:
        return

    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    is_four_col_demographics = (num_cols == 4)

    for r_idx, row in enumerate(rows):
        for c_idx, cell_value in enumerate(row):
            if c_idx < num_cols:
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                apply_cell_padding(cell, top=100, bottom=100, left=130, right=130)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                
                # Header row styling or 4-column key columns styling
                if r_idx == 0:
                    shading_elm = parse_xml(r'<w:shd {} w:fill="F2F4F7"/>'.format(nsdecls('w')))
                    cell._tc.get_or_add_tcPr().append(shading_elm)
                    apply_inline_formatting(p, cell_value, default_size_pt=10.5)
                    for r in p.runs:
                        r.bold = True
                elif is_four_col_demographics and c_idx in (0, 2):
                    apply_inline_formatting(p, cell_value, default_size_pt=10)
                    for r in p.runs:
                        r.bold = True
                else:
                    apply_inline_formatting(p, cell_value, default_size_pt=10)

    # Add small spacing after table
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(3)
    spacer.paragraph_format.space_after = Pt(3)


def build_docx_from_markdown(md_content: str) -> DocumentType:
    """Builds a python-docx Document object from structured Markdown."""
    doc = Document()

    # Set 0.75 inch page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    lines = md_content.splitlines()
    table_buffer = []
    in_table = False
    in_header_block = True  # Top header block before the first '---'

    for line in lines:
        stripped = line.strip()

        # 1. Accumulate Markdown Table Lines
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_buffer.append(stripped)
            continue
        elif in_table:
            render_markdown_table(doc, table_buffer)
            table_buffer = []
            in_table = False

        if not stripped:
            continue

        # 2. Horizontal Divider Line (--- or ***) -> Marks end of top header block
        if re.match(r"^(\-{3,}|\_{3,}|\*{3,})$", stripped):
            in_header_block = False  # Everything below this line is left-aligned body
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'D0D5DD')
            pBdr.append(bottom)
            p._element.get_or_add_pPr().append(pBdr)
            continue

        # 3. Level 1 Heading (# Title)
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            if in_header_block:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[2:].strip())
            run.font.name = FONT_NAME
            run.font.size = Pt(13.5)
            run.font.color.rgb = BLACK_COLOR
            run.bold = True

        # 4. Level 2 Heading (## Section)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            if in_header_block:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[3:].strip())
            run.font.name = FONT_NAME
            run.font.size = Pt(12)
            run.font.color.rgb = BLACK_COLOR
            run.bold = True

        # 5. Level 3 Heading (### Sub-section)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            if in_header_block:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[4:].strip())
            run.font.name = FONT_NAME
            run.font.size = Pt(11)
            run.font.color.rgb = BLACK_COLOR
            run.bold = True

        # 6. Bullet Points (* item, - item, • item)
        elif stripped.startswith(("- ", "* ", "• ")):
            content = re.sub(r"^[-*•]\s+", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            apply_inline_formatting(p, content)

        # 7. Numbered Lists (1. item)
        elif re.match(r"^\d+\.\s+", stripped):
            content = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            apply_inline_formatting(p, content)

        # 8. Regular Body Paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            if in_header_block:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            apply_inline_formatting(p, stripped)

    if in_table:
        render_markdown_table(doc, table_buffer)

    return doc


def convert_markdown_to_docx_bytes(md_content: str) -> bytes:
    """Renders Markdown to in-memory DOCX bytes."""
    doc = build_docx_from_markdown(md_content)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def parse_patient_metadata_from_md(md_content: str) -> dict:
    """Extracts basic patient info from standard discharge summary markdown table for list views."""
    meta = {
        "patient_name": "",
        "age": "",
        "sex": "",
        "uhid_no": "",
        "date_of_admission": "",
        "date_of_discharge": "",
    }
    match_name = re.search(r"\|\s*(?:\*\*)?Patient\s*Name(?:\*\*)?\s*\|\s*([^|\n]+)\|", md_content, re.IGNORECASE)
    if match_name:
        meta["patient_name"] = match_name.group(1).strip()
    
    match_age_sex = re.search(r"\|\s*(?:\*\*)?Age\s*/\s*Sex(?:\*\*)?\s*\|\s*([^|\n]+)\|", md_content, re.IGNORECASE)
    if match_age_sex:
        val = match_age_sex.group(1).strip()
        parts = [p.strip() for p in val.split("/") if p.strip()]
        if len(parts) >= 1:
            meta["age"] = parts[0]
        if len(parts) >= 2:
            meta["sex"] = parts[1]

    match_uhid = re.search(r"\|\s*(?:\*\*)?UHID(?:\s*No\.?)?(?:\*\*)?\s*\|\s*([^|\n]+)\|", md_content, re.IGNORECASE)
    if match_uhid:
        meta["uhid_no"] = match_uhid.group(1).strip()

    match_doa = re.search(r"\|\s*(?:\*\*)?Date\s*of\s*Admission(?:\*\*)?\s*\|\s*([^|\n]+)\|", md_content, re.IGNORECASE)
    if match_doa:
        meta["date_of_admission"] = match_doa.group(1).strip()

    match_dod = re.search(r"\|\s*(?:\*\*)?Date\s*of\s*Discharge(?:\*\*)?\s*\|\s*([^|\n]+)\|", md_content, re.IGNORECASE)
    if match_dod:
        meta["date_of_discharge"] = match_dod.group(1).strip()

    return meta

